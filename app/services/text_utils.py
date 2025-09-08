# app/services/text_utils.py
import os
import json
import csv
import re
from typing import List, Any
from pathlib import Path
from PyPDF2 import PdfReader
import docx
import openpyxl
from PIL import Image
import pytesseract
import cv2
import numpy as np
# OCR check
try:
    pytesseract.get_tesseract_version()
    _OCR_AVAILABLE = True
except Exception:
    _OCR_AVAILABLE = False

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff", ".jfif",".avif"}
DOCX_EXTS = {".docx"}
XLSX_EXTS = {".xlsx"}
TEXT_EXTS = {".txt", ".md", ".log"}
CSV_EXTS = {".csv"}
JSON_EXTS = {".json"}
PDF_EXTS = {".pdf"}
LEGACY_DOC_EXTS = {".doc"}
LEGACY_XLS_EXTS = {".xls"}
SVG_EXTS = {".svg"}

def _flatten_json(data: Any, prefix: str = "") -> List[str]:
    out = []
    if isinstance(data, dict):
        for k, v in data.items():
            key = f"{prefix}.{k}" if prefix else k
            out.extend(_flatten_json(v, key))
    elif isinstance(data, list):
        for i, v in enumerate(data):
            key = f"{prefix}[{i}]"
            out.extend(_flatten_json(v, key))
    else:
        out.append(f"{prefix}: {data}")
    return out

def _extract_from_pdf(path: str) -> str:
    text = ""
    try:
        reader = PdfReader(path)
        for p in reader.pages:
            text += p.extract_text() or ""
    except Exception:
        return ""
    return text

def _extract_from_docx(path: str) -> str:
    try:
        d = docx.Document(path)
        return "\n".join(para.text for para in d.paragraphs if para.text)
    except Exception:
        return ""

def _extract_from_xlsx(path: str) -> str:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception:
        return ""

def _extract_from_txt(path: str) -> str:
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        return ""

def _extract_from_csv(path: str) -> str:
    try:
        parts = []
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            for row in reader:
                if row:
                    parts.append(", ".join(row))
        return "\n".join(parts)
    except Exception:
        return ""

def _extract_from_json(path: str) -> str:
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            data = json.load(f)
        return "\n".join(_flatten_json(data))
    except Exception:
        return ""

def _extract_from_svg(path: str) -> str:
    try:
        raw = _extract_from_txt(path)
        return re.sub(r"<[^>]+>", " ", raw)
    except Exception:
        return ""

def _extract_from_image(path: str) -> str:
    if not _OCR_AVAILABLE:
        return ""
    try:

        img = Image.open(path)
        cv_img = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2GRAY)

        # Denoise and threshold
        cv_img = cv2.medianBlur(cv_img, 3)
        _, thresh = cv2.threshold(cv_img, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)

        # OCR
        text = pytesseract.image_to_string(thresh, config="--psm 6")
        print(f"[OCR DEBUG] Extracted text: {text!r}")
        return text.strip()
    except Exception as e:
        print(f"[OCR ERROR] {e}")
        return ""


def extract_text_from_file(path: str, content_type: str) -> str:
    ext = Path(path).suffix.lower()
    try:
        if ext in PDF_EXTS:
            return _extract_from_pdf(path)
        if ext in DOCX_EXTS:
            return _extract_from_docx(path)
        if ext in XLSX_EXTS:
            return _extract_from_xlsx(path)
        if ext in TEXT_EXTS:
            return _extract_from_txt(path)
        if ext in CSV_EXTS:
            return _extract_from_csv(path)
        if ext in JSON_EXTS:
            return _extract_from_json(path)
        if ext in SVG_EXTS:
            return _extract_from_svg(path)
        if ext in IMAGE_EXTS:
            return _extract_from_image(path)

        if ext in LEGACY_DOC_EXTS:
            try:
                return _extract_from_docx(path)
            except Exception:
                return ""
        if ext in LEGACY_XLS_EXTS:
            return _extract_from_txt(path)

        return _extract_from_txt(path)
    except Exception:
        return ""

def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> List[str]:
    if not text:
        return []
    chunks = []
    start = 0
    length = len(text)
    while start < length:
        end = min(start + chunk_size, length)
        chunk = text[start:end]
        chunks.append(chunk.strip())
        if end == length:
            break
        start = max(0, end - overlap)
    return [c for c in chunks if c]
